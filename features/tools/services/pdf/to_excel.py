"""PDF to Excel — extract tables using PyMuPDF + openpyxl."""
import fitz
import openpyxl

from features.tools.models import ToolJob
from features.tools.services import file_handler
from features.tools.services.exceptions import ToolError


def _extract_tables(doc: fitz.Document) -> list[list[list[str]]]:
    tables = []
    for page in doc:
        try:
            page_tables = page.find_tables()
            for tbl in page_tables.tables:
                rows = []
                for row in tbl.extract():
                    rows.append([str(cell or "") for cell in row])
                if rows:
                    tables.append(rows)
        except Exception:
            pass
    return tables


def _fallback_text_rows(doc: fitz.Document) -> list[list[str]]:
    rows = []
    for i, page in enumerate(doc):
        lines = page.get_text().splitlines()
        for line in lines:
            if line.strip():
                rows.append([f"Page {i+1}", line.strip()])
    return rows


def run(job: ToolJob) -> dict:
    src_rel = job.input_files[0]
    src = file_handler.to_absolute(src_rel)
    out = file_handler.output_path(job, "extracted.xlsx")

    doc = fitz.open(str(src))
    tables = _extract_tables(doc)
    wb = openpyxl.Workbook()

    if tables:
        for idx, table in enumerate(tables):
            ws = wb.create_sheet(title=f"Table {idx + 1}") if idx > 0 else wb.active
            ws.title = f"Table {idx + 1}"
            for row in table:
                ws.append(row)
    else:
        ws = wb.active
        ws.title = "Text Content"
        rows = _fallback_text_rows(doc)
        ws.append(["Page", "Text"])
        for row in rows:
            ws.append(row)

    doc.close()
    wb.save(str(out))
    out_rel = file_handler.to_relative(out)
    return {
        "output_file": out_rel,
        "compression_stat": {"tables_found": len(tables)},
        "preview_file": "",
    }
